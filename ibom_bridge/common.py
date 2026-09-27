"""Shared pieces: flattening an I-BOM into trackable assets, reading exports, findings."""
import csv
import io
import json
import re
from datetime import date
from pathlib import Path

MAPPINGS = Path(__file__).resolve().parent / "mappings"

# Lifecycle order used to decide whether one system is ahead of or behind another.
FLOW = ["planned", "quoted", "ordered", "backordered", "in-transit", "received", "staged",
        "burn-in", "deployed", "in-service"]
RECEIVED_OR_LATER = set(FLOW[FLOW.index("received"):]) | {
    "degraded", "failed", "in-repair", "rma", "spare", "decommissioned", "disposed"}
ACTIVE = {"deployed", "in-service", "degraded", "burn-in", "discovered-unmanaged"}
RETIRED = {"decommissioned", "disposed"}

SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}


def load_json(name):
    return json.loads((MAPPINGS / name).read_text())


def norm(value):
    """Loose comparison key: case-, space- and punctuation-insensitive."""
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def num(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "").replace("$", "")
    if s.endswith("-"):  # SAP writes negatives as 123.00-
        s = "-" + s[:-1]
    try:
        return float(s)
    except ValueError:
        return None


def status_of(component, unit=None):
    if unit and unit.get("lifecycleStatus"):
        return unit["lifecycleStatus"]
    return (component.get("lifecycle") or {}).get("status", "planned")


def event_date(component, statuses):
    """Date (YYYY-MM-DD) of the first lifecycle event that moved the line into one of `statuses`."""
    for e in (component.get("lifecycle") or {}).get("events", []):
        if e.get("type") == "status-change" and e.get("to") in statuses:
            return e["at"][:10]
    return None


def in_service_date(component):
    dates = (component.get("lifecycle") or {}).get("dates", {})
    return (dates.get("inService") or dates.get("installed")
            or event_date(component, {"deployed", "in-service"}))


class Doc:
    """Read-only helpers over an I-BOM document."""

    def __init__(self, doc):
        self.doc = doc
        self.orgs = {o["bom-ref"]: o for o in doc.get("organizations", [])}
        self.locs = {loc["bom-ref"]: loc for loc in doc.get("locations", [])}
        self.comps = {c["bom-ref"]: c for c in doc.get("components", [])}

    @classmethod
    def load(cls, path):
        return cls(json.loads(Path(path).read_text()))

    def org_name(self, ref):
        return (self.orgs.get(ref) or {}).get("name", ref or "")

    def loc_name(self, ref):
        return (self.locs.get(ref) or {}).get("name", ref or "")

    def vendor_id(self, component):
        for s in (component.get("supply") or {}).get("sources", []):
            vid = ((self.orgs.get(s.get("supplier")) or {}).get("identifiers") or {}).get("erpVendorId")
            if vid:
                return vid
        return (component.get("identifiers") or {}).get("erp", {}).get("vendorId", "")

    def assets(self):
        """One entry per trackable item.

        A line with serialized `units` yields one asset per unit, a quantity-1 line
        yields one asset, and a non-serialized quantity line (10 optics) yields one
        aggregate asset carrying the quantity.
        """
        for c in self.doc.get("components", []):
            units = c.get("units") or []
            if units:
                for i, u in enumerate(units):
                    yield Asset(self, c, u, i)
            else:
                yield Asset(self, c, None, None)


class Asset:
    def __init__(self, doc, component, unit, index):
        self.d = doc
        self.c = component
        self.u = unit
        self.index = index

    @property
    def ref(self):
        """Stable id for this asset, used as ServiceNow correlation_id and in reports."""
        return f"{self.c['bom-ref']}#{self.serial}" if self.u else self.c["bom-ref"]

    @property
    def serial(self):
        return (self.u or {}).get("serialNumber") or self.c.get("serialNumber")

    @property
    def asset_tag(self):
        return (self.u or {}).get("assetTag") or self.c.get("assetTag")

    @property
    def quantity(self):
        return 1 if self.u else self.c.get("quantity", 1)

    @property
    def serialized(self):
        return bool(self.serial)

    @property
    def status(self):
        return status_of(self.c, self.u)

    @property
    def ids(self):
        """Identifiers for this asset: unit-level values win over line-level ones."""
        line = self.c.get("identifiers") or {}
        unit = (self.u or {}).get("identifiers") or {}
        out = {}
        for k in set(line) | set(unit):
            a, b = line.get(k), unit.get(k)
            out[k] = {**a, **b} if isinstance(a, dict) and isinstance(b, dict) else (b if b is not None else a)
        return out

    @property
    def hostname(self):
        return (self.ids.get("other") or {}).get("hostname")

    @property
    def cloud_id(self):
        return (self.ids.get("cloud") or {}).get("resourceId")

    @property
    def location(self):
        return ((self.u or {}).get("placement") or self.c.get("placement") or {}).get("location")

    def set_ids(self, system, values):
        """Write identifiers back onto the unit (if serialized) or the line."""
        target = self.u if self.u is not None else self.c
        block = target.setdefault("identifiers", {}).setdefault(system, {})
        changed = False
        for k, v in values.items():
            if v and block.get(k) != v:
                block[k] = v
                changed = True
        return changed


def read_table(path):
    """Rows (list of dicts) from CSV, XLSX, or JSON (a list, or ServiceNow's {"result": [...]})."""
    p = Path(path)
    if p.suffix.lower() == ".json":
        data = json.loads(p.read_text())
        if isinstance(data, dict):
            data = data.get("result", data.get("records", data.get("items", [])))
        return [dict(r) for r in data]
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        head = next(i for i, r in enumerate(rows) if sum(v not in (None, "") for v in r) >= 2)
        cols = [str(v).strip() if v is not None else "" for v in rows[head]]
        return [{cols[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(cols) and cols[i]}
                for r in rows[head + 1:] if any(v not in (None, "") for v in r)]
    text = p.read_text(encoding="utf-8-sig")
    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t|")
    return [{(k or "").strip(): (v or "").strip() for k, v in r.items()}
            for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


def normalize_rows(rows, aliases):
    """Rename columns to canonical keys using {canonical: [alias, ...]} (case/punctuation-insensitive)."""
    lookup = {}
    for key, names in aliases.items():
        for n in [key, *names]:
            lookup.setdefault(norm(n), key)
    out = []
    for r in rows:
        row = {}
        for k, v in r.items():
            key = lookup.get(norm(k))
            if key and (key not in row or row[key] in ("", None)):
                row[key] = v
        out.append(row)
    return out


def write_csv(path, rows, columns):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore", lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in columns})


def finding(system, kind, severity, message, action=None, asset=None, bom_ref=None, record=None,
            field=None, bom=None, other=None):
    f = {"system": system, "kind": kind, "severity": severity}
    if asset is not None:
        f["bomRef"] = asset.c["bom-ref"]
        if asset.serial:
            f["serial"] = asset.serial
    elif bom_ref:
        f["bomRef"] = bom_ref
    if record:
        f["record"] = record
    if field:
        f["field"] = field
    if bom is not None:
        f["bom"] = bom
    if other is not None:
        f["other"] = other
    f["message"] = message
    if action:
        f["action"] = action
    return f


def sort_findings(findings):
    return sorted(findings, key=lambda f: (SEVERITY_ORDER[f["severity"]], f["system"], f["kind"],
                                           f.get("bomRef", ""), f.get("record", "")))


def today():
    return date.today().isoformat()
